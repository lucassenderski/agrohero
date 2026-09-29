#!/usr/bin/env python3
"""Converte o relatorio da AE III (Markdown) em DOCX formatado (padrao ABNT simplificado)."""
import re
import sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ENTRADA = sys.argv[1] if len(sys.argv) > 1 else 'docs/entrega/AE3-Relatorio.md'
SAIDA = sys.argv[2] if len(sys.argv) > 2 else 'AE3-Relatorio.docx'
FONTE = 'Times New Roman'
MONO = 'Consolas'


def definir_fonte(estilo, nome, tamanho, negrito=False, italico=False):
    estilo.font.name = nome
    estilo.font.size = Pt(tamanho)
    estilo.font.bold = negrito
    estilo.font.italic = italico
    rpr = estilo.element.get_or_add_rPr()
    fontes = rpr.find(qn('w:rFonts'))
    if fontes is None:
        fontes = OxmlElement('w:rFonts')
        rpr.append(fontes)
    for atributo in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fontes.set(qn(atributo), nome)


def formatar_inline(paragrafo, texto):
    """Aplica negrito/itálico/código do Markdown em um parágrafo já existente."""
    texto = texto.replace('\\|', '|')
    padrao = re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)', re.S)
    for parte in padrao.split(texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**') and len(parte) > 4:
            r = paragrafo.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith('`') and parte.endswith('`') and len(parte) > 2:
            r = paragrafo.add_run(parte[1:-1]); r.font.name = MONO
            r.font.size = Pt(10)
        elif parte.startswith('*') and parte.endswith('*') and len(parte) > 2:
            r = paragrafo.add_run(parte[1:-1]); r.italic = True
        else:
            paragrafo.add_run(parte)


def sombrear(celula, cor='EEEEEE'):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), cor)
    tcpr.append(shd)


def separar_cabecalho(linha):
    return [c.strip() for c in linha.strip().strip('|').split('|')]


def eh_separador_tabela(linha):
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', linha.strip()))


def montar_capa(doc, meta):
    for _ in range(2):
        doc.add_paragraph()
    for texto, tamanho, negrito in [
        (meta['instituicao'], 12, True),
        (meta['curso'], 12, False),
        (meta['academico'], 12, False),
        (meta['ru'], 12, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto); r.bold = negrito

    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ATIVIDADE EXTENSIONISTA III'); r.bold = True; r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Tecnologia Aplicada à Inclusão Digital — Análise')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Evolução e Gerenciamento do Projeto AgroHero')
    r.bold = True; r.font.size = Pt(13)

    for _ in range(6):
        doc.add_paragraph()

    for texto in [meta['cidade'], meta['ano']]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(texto)

    doc.add_page_break()


def main():
    meta = {'instituicao': '[NOME DA INSTITUIÇÃO DE ENSINO]',
            'curso': '[CURSO]',
            'academico': 'Lucas Senderski',
            'ru': 'RU: 4758862',
            'cidade': '[CIDADE]',
            'ano': '2026'}
    linhas = open(ENTRADA, encoding='utf-8').read().split('\n')

    doc = Document()
    secao = doc.sections[0]
    secao.top_margin = Cm(3); secao.left_margin = Cm(3)
    secao.right_margin = Cm(2); secao.bottom_margin = Cm(2)

    normal = doc.styles['Normal']
    definir_fonte(normal, FONTE, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nivel, tam in ((1, 14), (2, 13), (3, 12), (4, 12)):
        e = doc.styles[f'Heading {nivel}']
        definir_fonte(e, FONTE, tam, negrito=True)
        e.font.color.rgb = RGBColor(0x1B, 0x1B, 0x1B)
        e.paragraph_format.space_before = Pt(12)
        e.paragraph_format.space_after = Pt(6)
        e.paragraph_format.line_spacing = 1.5

    montar_capa(doc, meta)

    i = 0
    primeiro_h1 = True
    while i < len(linhas):
        linha = linhas[i]

        if linha.strip() in ('---', '***', '___'):
            i += 1
            continue

        if not linha.strip():
            i += 1
            continue

        if linha.startswith('```'):
            bloco = []
            i += 1
            while i < len(linhas) and not linhas[i].startswith('```'):
                bloco.append(linhas[i]); i += 1
            i += 1
            for b in bloco:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(b if b else ' '); r.font.name = MONO; r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        m = re.match(r'^(#{1,6})\s+(.*)$', linha)
        if m:
            nivel = len(m.group(1)); texto = m.group(2).strip()
            if nivel == 1:
                if not primeiro_h1:
                    doc.add_page_break()
                primeiro_h1 = False
                texto = re.sub(r'\s+#+$', '', texto)
            p = doc.add_heading(level=min(nivel, 4))
            p.text = ''
            formatar_inline(p, texto)
            for r in p.runs:
                r.font.name = FONTE
            i += 1
            continue

        if linha.strip().startswith('|') and i + 1 < len(linhas) and eh_separador_tabela(linhas[i + 1]):
            cabecalho = separar_cabecalho(linha)
            i += 2
            corpo = []
            while i < len(linhas) and linhas[i].strip().startswith('|'):
                corpo.append(separar_cabecalho(linhas[i])); i += 1
            tabela = doc.add_table(rows=1, cols=len(cabecalho))
            tabela.style = 'Table Grid'
            tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, c in enumerate(cabecalho):
                cel = tabela.rows[0].cells[j]
                cel.text = ''
                formatar_inline(cel.paragraphs[0], c)
                for r in cel.paragraphs[0].runs:
                    r.bold = True; r.font.size = Pt(10.5); r.font.name = FONTE
                sombrear(cel)
            for linha_corpo in corpo:
                celulas = tabela.add_row().cells
                for j, c in enumerate(linha_corpo[:len(cabecalho)]):
                    celulas[j].text = ''
                    p = celulas[j].paragraphs[0]
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(2)
                    formatar_inline(p, c)
                    for r in p.runs:
                        r.font.size = Pt(10.5); r.font.name = FONTE
            doc.add_paragraph()
            continue

        if linha.strip().startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            formatar_inline(p, linha.strip()[2:])
            for r in p.runs:
                r.font.size = Pt(11); r.italic = True
            i += 1
            continue

        m = re.match(r'^(\s*)[-*]\s+(.*)$', linha)
        if m:
            recuo = len(m.group(1)) // 2
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(1.25 + recuo * 0.6)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            formatar_inline(p, m.group(2))
            i += 1
            continue

        m = re.match(r'^\s*(\d+)\.\s+(.*)$', linha)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run(f'{m.group(1)}. ')
            formatar_inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        formatar_inline(p, linha.strip())
        i += 1

    doc.save(SAIDA)
    print(f'Gerado: {SAIDA}')


if __name__ == '__main__':
    main()
